#!/usr/bin/env python3
"""Create matching editable and PDF classroom links slides from one layout."""
from pathlib import Path
import io

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

OUT = Path('site_assets/downloads')
WIDTH, HEIGHT = 960, 540
FG, ACCENT, MUTED = '#123044', '#0b6e6a', '#526776'
PORTAL = 'https://coursemgr.advisorssrcdemo.src.surf-hosted.nl/'
COMMANDS = 'https://supersurf.rgoswami.me/commands'
SETUP = 'https://supersurf.rgoswami.me/setup'
SLIDES = 'https://supersurf.rgoswami.me/downloads/introduction-slides.zip'


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = 'Course links and Slurm tips'
    prs.core_properties.author = 'SURF course materials'
    pdf = canvas.Canvas(str(OUT / 'course-links.pdf'), pagesize=(WIDTH, HEIGHT))
    pdf.setTitle('Course links and Slurm tips')
    pdf.setAuthor('SURF course materials')
    slide = None

    def new_slide():
        nonlocal slide
        slide = prs.slides.add_slide(prs.slide_layouts[6])

    def text(value, x, y, w, size=20, color=FG, bold=False, url=None, mono=False):
        lines = value.splitlines()
        box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt((size * 1.35) * len(lines) + 8))
        frame = box.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.2
            run = paragraph.add_run()
            run.text = line
            run.font.name = 'Courier New' if mono else 'Arial'
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = RGBColor.from_string(color.lstrip('#'))
            if url:
                run.hyperlink.address = url
            font = 'Courier' if mono else ('Helvetica-Bold' if bold else 'Helvetica')
            pdf.setFont(font, size)
            pdf.setFillColor(HexColor(color))
            baseline = HEIGHT - y - size - i * size * 1.2
            pdf.drawString(x, baseline, line)
            if url:
                pdf.linkURL(url, (x, baseline - 2, x + w, baseline + size), relative=0, thickness=0)

    new_slide()
    text('Course links', 64, 45, 800, 40, bold=True)
    text('Open the command page and keep it beside your editor.', 64, 103, 820, 20, MUTED)
    text('Copy commands', 64, 155, 620, 24, bold=True)
    text('supersurf.rgoswami.me/commands', 64, 190, 630, 24, ACCENT, url=COMMANDS)
    text('Course portal', 64, 246, 650, 21, bold=True)
    text('coursemgr.advisorssrcdemo.src.surf-hosted.nl', 64, 278, 650, 17, ACCENT, url=PORTAL)
    text('Setup guide', 64, 324, 620, 21, bold=True)
    text('supersurf.rgoswami.me/setup', 64, 356, 630, 21, ACCENT, url=SETUP)
    text('Download the presentation PDFs', 64, 407, 640, 21, ACCENT, url=SLIDES)
    qr = qrcode.make(COMMANDS)
    buffer = io.BytesIO()
    qr.save(buffer, format='PNG')
    buffer.seek(0)
    slide.shapes.add_picture(buffer, Pt(742), Pt(156), width=Pt(154), height=Pt(154))
    buffer.seek(0)
    pdf.drawImage(ImageReader(buffer), 742, HEIGHT - 156 - 154, 154, 154)
    text('Scan for commands', 735, 324, 180, 16, MUTED)
    text('Use your own private account-setup link to choose your password.', 64, 488, 830, 15, MUTED)
    pdf.showPage()

    new_slide()
    text('Jobs and output', 64, 45, 820, 40, bold=True)
    text('Create the directory before submitting. Keep the printed job number.', 64, 103, 840, 19, MUTED)
    command = '''mkdir -p logs
sbatch --partition=rome --cpus-per-task=1 \\
  --job-name=python-check \\
  --output='logs/%x-%j.out' \\
  --error='logs/%x-%j.err' python.sh'''
    text(command, 64, 155, 850, 22, mono=True)
    text('%x = job name    %j = job ID    %A_%a = array job and task', 64, 320, 840, 19)
    text('squeue --me\nsacct -j JOBID --format=JobID,State,ExitCode\nseff JOBID', 64, 368, 840, 19, mono=True)
    text('Choose 1 core for VS Code. Stop its session when finished. No email notifications.', 64, 488, 850, 15, MUTED)
    pdf.showPage()
    pdf.save()
    prs.save(OUT / 'course-links.pptx')
    print('LINKS_DECK_CREATED slides=2 formats=pptx,pdf')


if __name__ == '__main__':
    main()
